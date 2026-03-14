import os
import sys
from typing import List

import tiktoken
from rich.console import Console
from rich.panel import Panel

from pipeline.models import Document

console = Console()

SUPPORTED_EXTENSIONS = {".md", ".txt", ".pdf", ".docx"}


def _extract_text(filepath: str, ext: str) -> str:
    if ext in (".md", ".txt"):
        with open(filepath, encoding="utf-8", errors="replace") as f:
            return f.read()

    if ext == ".pdf":
        import pdfplumber
        pages = []
        with pdfplumber.open(filepath) as pdf:
            for page in pdf.pages:
                text = page.extract_text()
                if text:
                    pages.append(text)
        return "\n".join(pages)

    if ext == ".docx":
        import docx
        doc = docx.Document(filepath)
        return "\n".join(paragraph.text for paragraph in doc.paragraphs)

    raise ValueError(f"Unsupported extension: {ext}")


def extract(folder_path: str) -> List[Document]:
    encoding = tiktoken.get_encoding("cl100k_base")

    entries = [
        e for e in os.scandir(folder_path)
        if e.is_file() and os.path.splitext(e.name)[1].lower() in SUPPORTED_EXTENSIONS
    ]

    if not entries:
        console.print(Panel(
            f"No supported files found in [bold]{folder_path}[/bold].\n"
            f"Supported formats: {', '.join(sorted(SUPPORTED_EXTENSIONS))}",
            title="[red]Error[/red]",
            border_style="red",
        ))
        sys.exit(1)

    documents: List[Document] = []

    for entry in entries:
        ext = os.path.splitext(entry.name)[1].lower()
        try:
            content = _extract_text(entry.path, ext)
            token_count = len(encoding.encode(content))
            documents.append(Document(
                filename=entry.name,
                content=content,
                token_count=token_count,
            ))
        except Exception as exc:
            console.print(
                f"[yellow]Warning:[/yellow] Skipping [bold]{entry.name}[/bold] — {exc}"
            )

    if not documents:
        console.print(Panel(
            "All files failed extraction. No documents could be processed.",
            title="[red]Error[/red]",
            border_style="red",
        ))
        sys.exit(1)

    return documents
